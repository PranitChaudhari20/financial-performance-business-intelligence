"""
06_verify.py  -  do Python, MySQL and Excel agree?
==================================================
Author: Pranit Ajay Chaudhari

The project's credibility rests on one claim: the same figures come out of all three tools. This
script tests that claim against the actual files and fails loudly if it is false.

Exit codes
    0  every check passed
    1  at least one figure disagrees
    2  the Excel workbook has no cached values, so its checks COULD NOT BE RUN
       (run python/07_recalc_excel.py, or open the workbook in Excel and save)

Run:  python python/06_verify.py
"""
import sys
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
CSV = ROOT / "outputs" / "csv"
Q = ROOT / "sql" / "query_results"
XL = ROOT / "excel" / "Financial_Performance_Analytics.xlsx"

OK, BAD = [], []


def num(x):
    try:
        return float(str(x).replace(",", ""))
    except (TypeError, ValueError):
        return None


def check(name, a, b, tol=0.01):
    na, nb = num(a), num(b)
    same = (str(a).strip() == str(b).strip()) if na is None or nb is None else abs(na - nb) <= tol
    (OK if same else BAD).append(name)
    print(f"  {'PASS ' if same else 'FAIL '} {name:<58} {a} | {b}")


print("CROSS-TOOL VERIFICATION\n" + "=" * 94)

kpi = pd.read_csv(CSV / "kpi_summary.csv").set_index("Metric").Value

# ---------------------------------------------------------- Python vs SQL
print("\nPython vs MySQL")
q1 = pd.read_csv(Q / "q01_business_kpis.csv").iloc[0]
for label, sql_col in [("Order lines", "order_lines"), ("Orders", "orders"),
                       ("Units sold", "units_sold"), ("Distinct products", "distinct_products"),
                       ("Distinct customers", "distinct_customers"),
                       ("Gross sales (list value)", "gross_sales"),
                       ("Discount given", "discount_given"),
                       ("Discount rate (% of gross sales)", "discount_rate_pct"),
                       ("Net revenue", "net_revenue"),
                       ("Implied cost of sales", "implied_cost"), ("Profit", "profit"),
                       ("Net margin (% of net revenue)", "net_margin_pct"),
                       ("Cost-to-revenue ratio (%)", "cost_to_revenue_pct"),
                       ("Average order value", "avg_order_value"),
                       ("Profit per order", "profit_per_order"),
                       ("Loss-making order lines", "loss_making_lines"),
                       ("Loss-making order lines (%)", "loss_making_lines_pct")]:
    check(f"KPI {label}", kpi[label], q1[sql_col])

py_yr = pd.read_csv(CSV / "yearly_performance.csv").set_index("Year")
q2 = pd.read_csv(Q / "q02_yearly_performance_and_growth.csv").set_index("year")
for y in py_yr.index:
    check(f"{y} revenue", py_yr.loc[y, "Revenue"], q2.loc[y, "revenue"])
    check(f"{y} profit", py_yr.loc[y, "Profit"], q2.loc[y, "profit"])
    check(f"{y} margin %", py_yr.loc[y, "MarginPct"], q2.loc[y, "margin_pct"])

py_cat = pd.read_csv(CSV / "category_performance.csv").set_index("Category")
q6 = pd.read_csv(Q / "q06_category_performance.csv").set_index("category")
for c in py_cat.index:
    check(f"category {c}: revenue", py_cat.loc[c, "Revenue"], q6.loc[c, "revenue"])
    check(f"category {c}: profit", py_cat.loc[c, "Profit"], q6.loc[c, "profit"])
    check(f"category {c}: margin %", py_cat.loc[c, "MarginPct"], q6.loc[c, "margin_pct"])

py_sub = pd.read_csv(CSV / "subcategory_performance.csv").set_index("Sub-Category")
q7 = pd.read_csv(Q / "q07_sub_category_performance.csv").set_index("sub_category")
mism = sum(1 for s in py_sub.index
           if abs(float(py_sub.loc[s, "Profit"]) - float(q7.loc[s, "profit"])) > 0.01
           or abs(float(py_sub.loc[s, "MarginPct"]) - float(q7.loc[s, "margin_pct"])) > 0.01)
check(f"all {len(py_sub)} sub-categories: profit and margin", mism, 0)

py_reg = pd.read_csv(CSV / "regional_performance.csv").set_index("Region")
q8 = pd.read_csv(Q / "q08_regional_performance.csv").set_index("region")
for r_ in py_reg.index:
    check(f"region {r_}: revenue", py_reg.loc[r_, "Revenue"], q8.loc[r_, "revenue"])
    check(f"region {r_}: margin %", py_reg.loc[r_, "MarginPct"], q8.loc[r_, "margin_pct"])
    check(f"region {r_}: discount rate %", py_reg.loc[r_, "DiscountRatePct"], q8.loc[r_, "discount_rate_pct"])

py_db = pd.read_csv(CSV / "discount_band_performance.csv").set_index("DiscountBand")
q10 = pd.read_csv(Q / "q10_discount_band_versus_margin.csv").set_index("discount_band")
for b in py_db.index:
    check(f"discount band {b}: margin %", py_db.loc[b, "MarginPct"], q10.loc[b, "margin_pct"])
    check(f"discount band {b}: loss-line %", py_db.loc[b, "LossLineSharePct"], q10.loc[b, "loss_line_share_pct"])

py_st = pd.read_csv(CSV / "state_performance.csv")
py_loss = py_st[py_st.Profit < 0].set_index("State")
q9 = pd.read_csv(Q / "q09_states_that_lose_money.csv").set_index("state")
check("loss-making states: count", len(py_loss), len(q9))
check("loss-making states: total loss", round(py_loss.Profit.sum(), 2), round(q9.profit.sum(), 2))

q12 = pd.read_csv(Q / "q12_cost_structure.csv").iloc[0]
py_cost = pd.read_csv(CSV / "cost_structure.csv").set_index("Component").Amount
check("cost structure: gross sales", py_cost["Gross sales (list value)"], q12.gross_sales)
check("cost structure: net revenue", py_cost["= Net revenue"], q12.net_revenue)
check("cost structure: implied cost", abs(py_cost["less: implied cost of sales"]), q12.implied_cost_of_sales)
check("cost structure: profit", py_cost["= Profit"], q12.profit)

# --------------------------------------------------------- Python vs Excel
print("\nPython vs Excel (cached values)")
wb = load_workbook(XL, data_only=True)
ws = wb["KPI_Summary"]
xl = {}
for r_ in range(4, 40):
    a, b = ws.cell(r_, 1).value, ws.cell(r_, 2).value
    if a and b is not None:
        xl[a] = b

if "Net revenue" not in xl:
    print("  SKIPPED - the workbook has no cached formula values.")
    print("            openpyxl writes formulas but cannot evaluate them, so a freshly generated")
    print("            workbook reads back empty. Fix with either of:")
    print("              python python/07_recalc_excel.py     (needs LibreOffice on PATH)")
    print("              open excel/Financial_Performance_Analytics.xlsx in Excel and save")
    print("            then run this script again.")
    print("\n" + "=" * 94)
    print(f"  CHECKS PASSED : {len(OK)}")
    print(f"  CHECKS FAILED : {len(BAD)}")
    print("  EXCEL CHECKS  : NOT RUN (see above)")
    print("=" * 94)
    sys.exit(1 if BAD else 2)

check("Excel order lines", kpi["Order lines"], xl["Order lines"])
check("Excel orders", kpi["Orders"], xl["Orders"])
check("Excel units sold", kpi["Units sold"], xl["Units sold"])
check("Excel gross sales", kpi["Gross sales (list value)"], xl["Gross sales (list value)"])
check("Excel discount given", kpi["Discount given"], xl["Discount given"])
check("Excel discount rate %", kpi["Discount rate (% of gross sales)"], xl["Discount rate"] * 100)
check("Excel net revenue", kpi["Net revenue"], xl["Net revenue"])
check("Excel implied cost", kpi["Implied cost of sales"], xl["Implied cost of sales"])
check("Excel profit", kpi["Profit"], xl["Profit"])
check("Excel net margin %", kpi["Net margin (% of net revenue)"], xl["Net margin"] * 100)
check("Excel cost-to-revenue %", kpi["Cost-to-revenue ratio (%)"], xl["Cost-to-revenue ratio"] * 100)
check("Excel average order value", kpi["Average order value"], xl["Average order value"])
check("Excel profit per order", kpi["Profit per order"], xl["Profit per order"])
check("Excel loss-making lines", kpi["Loss-making order lines"], xl["Loss-making order lines"])
check("Excel loss-making lines %", kpi["Loss-making order lines (%)"],
      xl["Loss-making order lines (share)"] * 100)
check("Excel profit lost on those lines", kpi["Profit lost on loss-making lines"],
      xl["Profit lost on those lines"])

wd = wb["Discount_Analysis"]
for i, band in enumerate(py_db.index):
    check(f"Excel discount band {band}: margin %", py_db.loc[band, "MarginPct"],
          wd.cell(5 + i, 8).value * 100)

wc = wb["Cost_Analysis"]
check("Excel cost structure: gross sales", py_cost["Gross sales (list value)"], wc["B5"].value)
check("Excel cost structure: net revenue", py_cost["= Net revenue"], wc["B7"].value)
check("Excel cost structure: profit", py_cost["= Profit"], wc["B9"].value)

print("\n" + "=" * 94)
print(f"  CHECKS PASSED : {len(OK)}")
print(f"  CHECKS FAILED : {len(BAD)}")
for b in BAD:
    print(f"     - {b}")
print("=" * 94)
sys.exit(1 if BAD else 0)
