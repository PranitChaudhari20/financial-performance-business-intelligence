"""
05_build_excel.py  -  the analytical workbook
=============================================
Author: Pranit Ajay Chaudhari

Input : data/cleaned/superstore_clean.csv
Output: excel/Financial_Performance_Analytics.xlsx

Every reported figure is a LIVE FORMULA over the Data sheet, not a value pasted from Python. That
is the point: the workbook recomputes the analysis independently, so if Excel and Python disagree
the disagreement is visible rather than hidden. python/06_verify.py checks exactly that.

openpyxl writes formulas but cannot evaluate them, so a freshly generated workbook has no cached
values and previews blank. python/07_recalc_excel.py fixes that and run_all.sh calls it.

Run:  python python/05_build_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "Financial_Performance_Analytics.xlsx"
OUT.parent.mkdir(exist_ok=True)

AUTHOR = "Pranit Ajay Chaudhari"
FONT = "Arial"
INK, ACCENT, MUTED, GOOD, BAD = "FF0B0B0B", "FF2A78D6", "FF52514E", "FF1BAF7A", "FFD03B3B"
HDR = PatternFill("solid", fgColor="FFE8F0FB")
TITLE = PatternFill("solid", fgColor="FF2A78D6")
BAND = PatternFill("solid", fgColor="FF52514E")
BORDER = Border(bottom=Side(style="thin", color="FFD9D9D9"))
MONEY, PCT2, NUM2, INT = '#,##0.00', '0.00%', '#,##0.00', '#,##0'

d = pd.read_csv(ROOT / "data" / "cleaned" / "superstore_clean.csv", parse_dates=["OrderDate"])
print(f"  {len(d):,} order lines")

wb = Workbook()
wb.properties.creator = AUTHOR
wb.properties.lastModifiedBy = AUTHOR
wb.properties.title = "Financial Performance & Business Intelligence Analytics"


def title_block(ws, text, sub=""):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="FFFFFFFF")
    ws["A1"].fill = TITLE
    ws.row_dimensions[1].height = 24
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color=MUTED)


def head(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        c.fill = HDR
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


# ================================================================ DATA SHEET
COLS = ["Row ID", "Order ID", "OrderDate", "Segment", "Region", "State", "Category",
        "Sub-Category", "Quantity", "Discount", "DiscountBand", "GrossSales",
        "DiscountValue", "Revenue", "ImpliedCost", "Profit", "IsLossMaking",
        "IsFirstLineOfOrder", "Year", "YearMonth"]
data = d[COLS].copy()
data["OrderDate"] = data.OrderDate.dt.strftime("%Y-%m-%d")
wsD = wb.active
wsD.title = "Data"
title_block(wsD, "Cleaned order lines - the single source every formula reads",
            f"{len(data):,} rows. Produced by python/01_clean_data.py. Author: {AUTHOR}")
head(wsD, 4, COLS)
bool_i = COLS.index("IsLossMaking")
for row in data.itertuples(index=False):
    vals = list(row)
    vals[bool_i] = int(vals[bool_i])      # written as 1/0 so LibreOffice does not rewrite as =TRUE()
    wsD.append(vals)
LAST = 4 + len(data)
wsD.freeze_panes = "A5"
C = {c: get_column_letter(i + 1) for i, c in enumerate(COLS)}


def rng(col):
    return f"Data!${C[col]}$5:${C[col]}${LAST}"


print(f"  Data sheet rows 5..{LAST}")

# ========================================================== KPI SUMMARY ====
ws = wb.create_sheet("KPI_Summary")
title_block(ws, "Business KPIs", f"Every value is a live formula over the Data sheet. Author: {AUTHOR}")
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 62
KPI = [
    ("VOLUME", None, None, None),
    ("Order lines", f"=COUNTA({rng('Row ID')})", INT, "One row per product on an order"),
    ("Orders", f"=SUM({rng('IsFirstLineOfOrder')})", INT,
     "Distinct Order IDs, counted via the first-line flag rather than an array formula"),
    ("Units sold", f"=SUM({rng('Quantity')})", INT, None),
    ("REVENUE", None, None, None),
    ("Gross sales (list value)", f"=SUM({rng('GrossSales')})", MONEY, "Revenue / (1 - Discount)"),
    ("Discount given", f"=SUM({rng('DiscountValue')})", MONEY, "Gross sales - net revenue"),
    ("Discount rate", f"=SUM({rng('DiscountValue')})/SUM({rng('GrossSales')})", PCT2, "Share of list value given away"),
    ("Net revenue", f"=SUM({rng('Revenue')})", MONEY, "The dataset's Sales, already net of discount"),
    ("COST AND PROFIT", None, None, None),
    ("Implied cost of sales", f"=SUM({rng('ImpliedCost')})", MONEY, "Net revenue - profit. A RESIDUAL, not a reported cost line"),
    ("Profit", f"=SUM({rng('Profit')})", MONEY, "The dataset's own per-line profit"),
    ("Net margin", f"=SUM({rng('Profit')})/SUM({rng('Revenue')})", PCT2, "Profit / net revenue"),
    ("Cost-to-revenue ratio", f"=SUM({rng('ImpliedCost')})/SUM({rng('Revenue')})", PCT2, None),
    ("PER-ORDER ECONOMICS", None, None, None),
    ("Average order value", f"=SUM({rng('Revenue')})/B6", MONEY, "Net revenue / orders"),
    ("Profit per order", f"=SUM({rng('Profit')})/B6", MONEY, None),
    ("Average revenue per line", f"=SUM({rng('Revenue')})/B5", MONEY, None),
    ("Profit per unit", f"=SUM({rng('Profit')})/SUM({rng('Quantity')})", NUM2, None),
    ("LOSS EXPOSURE", None, None, None),
    ("Loss-making order lines", f"=SUM({rng('IsLossMaking')})", INT, "Lines where profit is below zero"),
    ("Loss-making order lines (share)", f"=SUM({rng('IsLossMaking')})/COUNTA({rng('Row ID')})", PCT2, None),
    ("Profit lost on those lines", f"=SUMIF({rng('Profit')},\"<0\")", MONEY, "What the business gives back"),
]
r_ = 4
ROWREF = {}
for lab, f, fmt, note in KPI:
    c = ws.cell(row=r_, column=1, value=lab)
    if f is None:
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFFFF")
        for col in (1, 2, 3):
            ws.cell(row=r_, column=col).fill = BAND
    else:
        ROWREF[lab] = r_
        c.font = Font(name=FONT, size=10, color=INK)
        v = ws.cell(row=r_, column=2, value=f)
        v.font = Font(name=FONT, size=10, bold=True, color=ACCENT)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="right")
        if note:
            n = ws.cell(row=r_, column=3, value=note)
            n.font = Font(name=FONT, size=9, italic=True, color=MUTED)
        for col in (1, 2, 3):
            ws.cell(row=r_, column=col).border = BORDER
    r_ += 1


def fin_block(ws, row, keys, key_range, label, extra_width=None):
    """A standard financial block computed with SUMIF over the Data sheet."""
    cols = ["Revenue", "% of revenue", "Implied cost", "Profit", "% of profit",
            "Net margin", "Cost to revenue", "Discount rate", "Loss lines %"]
    head(ws, row, [label] + cols, [extra_width or 22] + [14] * len(cols))
    first, last = row + 1, row + len(keys)
    for i, k in enumerate(keys):
        rr = row + 1 + i
        ws.cell(row=rr, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
        ws.cell(row=rr, column=2, value=f'=SUMIF({key_range},$A{rr},{rng("Revenue")})').number_format = MONEY
        ws.cell(row=rr, column=3, value=f'=B{rr}/SUM($B${first}:$B${last})').number_format = PCT2
        ws.cell(row=rr, column=4, value=f'=SUMIF({key_range},$A{rr},{rng("ImpliedCost")})').number_format = MONEY
        ws.cell(row=rr, column=5, value=f'=SUMIF({key_range},$A{rr},{rng("Profit")})').number_format = MONEY
        ws.cell(row=rr, column=6, value=f'=E{rr}/SUM($E${first}:$E${last})').number_format = PCT2
        ws.cell(row=rr, column=7, value=f'=E{rr}/B{rr}').number_format = PCT2
        ws.cell(row=rr, column=8, value=f'=D{rr}/B{rr}').number_format = PCT2
        ws.cell(row=rr, column=9,
                value=f'=SUMIF({key_range},$A{rr},{rng("DiscountValue")})'
                      f'/SUMIF({key_range},$A{rr},{rng("GrossSales")})').number_format = PCT2
        ws.cell(row=rr, column=10,
                value=f'=SUMIF({key_range},$A{rr},{rng("IsLossMaking")})'
                      f'/COUNTIF({key_range},$A{rr})').number_format = PCT2
    return last


# ======================================================= REVENUE ANALYSIS ==
ws2 = wb.create_sheet("Revenue_Analysis")
title_block(ws2, "Revenue analysis", "By year, category and region. All formulas.")
years = sorted(d.Year.unique().tolist())
end = fin_block(ws2, 4, years, rng("Year"), "Year", 12)
ws2.cell(row=end + 2, column=1, value="Year-on-year revenue growth").font = Font(name=FONT, size=10, bold=True)
for i, y in enumerate(years[1:], start=1):
    rr = 5 + i
    ws2.cell(row=end + 2 + i, column=1, value=y).font = Font(name=FONT, size=10)
    ws2.cell(row=end + 2 + i, column=2, value=f"=B{rr}/B{rr - 1}-1").number_format = PCT2

cats = sorted(d.Category.unique().tolist())
r2 = end + 2 + len(years) + 2
ws2.cell(row=r2, column=1, value="By category").font = Font(name=FONT, size=11, bold=True, color=ACCENT)
fin_block(ws2, r2 + 1, cats, rng("Category"), "Category", 20)
regs = sorted(d.Region.unique().tolist())
r3 = r2 + 1 + len(cats) + 3
ws2.cell(row=r3, column=1, value="By region").font = Font(name=FONT, size=11, bold=True, color=ACCENT)
fin_block(ws2, r3 + 1, regs, rng("Region"), "Region", 20)

# ========================================================= PROFITABILITY ===
ws3 = wb.create_sheet("Profitability")
title_block(ws3, "Profitability by sub-category",
            "Sorted by revenue. A negative net margin means the line loses money on every dollar sold.")
subs = d.groupby("Sub-Category").Revenue.sum().sort_values(ascending=False).index.tolist()
fin_block(ws3, 4, subs, rng("Sub-Category"), "Sub-Category", 22)

# ========================================================= COST ANALYSIS ===
ws4 = wb.create_sheet("Cost_Analysis")
title_block(ws4, "Cost structure",
            "From list value down to profit. The dataset has no operating expenses, so no operating "
            "margin or EBITDA can be computed and none is shown.")
head(ws4, 4, ["Component", "Amount", "% of gross sales", "% of net revenue"], [34, 18, 18, 18])
# NOTE: a label must never begin with "=", or Excel and LibreOffice parse it as a formula. An
# earlier version used "= Net revenue" and "= Profit" for the accounting look and both cells
# evaluated to #VALUE! and #NAME?. Subtotal rows are marked bold instead.
rowsc = [("Gross sales (list value)", f"=SUM({rng('GrossSales')})", False),
         ("less: discount given", f"=-SUM({rng('DiscountValue')})", False),
         ("Net revenue", f"=SUM({rng('Revenue')})", True),
         ("less: implied cost of sales", f"=-SUM({rng('ImpliedCost')})", False),
         ("Profit", f"=SUM({rng('Profit')})", True)]
for i, (lab, f, is_total) in enumerate(rowsc):
    rr = 5 + i
    ws4.cell(row=rr, column=1, value=lab).font = Font(name=FONT, size=10, bold=is_total)
    ws4.cell(row=rr, column=2, value=f).number_format = MONEY
    ws4.cell(row=rr, column=3, value=f"=B{rr}/$B$5").number_format = PCT2
    ws4.cell(row=rr, column=4, value=f"=B{rr}/$B$7").number_format = PCT2
ws4.cell(row=11, column=1, value="Cost-to-revenue ratio by year").font = Font(name=FONT, size=11, bold=True, color=ACCENT)
head(ws4, 12, ["Year", "Net revenue", "Implied cost", "Cost to revenue", "Net margin"], [12, 18, 18, 16, 16])
for i, y in enumerate(years):
    rr = 13 + i
    ws4.cell(row=rr, column=1, value=y)
    ws4.cell(row=rr, column=2, value=f'=SUMIF({rng("Year")},$A{rr},{rng("Revenue")})').number_format = MONEY
    ws4.cell(row=rr, column=3, value=f'=SUMIF({rng("Year")},$A{rr},{rng("ImpliedCost")})').number_format = MONEY
    ws4.cell(row=rr, column=4, value=f"=C{rr}/B{rr}").number_format = PCT2
    ws4.cell(row=rr, column=5, value=f'=SUMIF({rng("Year")},$A{rr},{rng("Profit")})/B{rr}').number_format = PCT2

# ============================================== DISCOUNT / CATEGORY ========
ws5 = wb.create_sheet("Discount_Analysis")
title_block(ws5, "Discount and margin",
            "The relationship that drives the loss-making lines. Discount bands are fixed and documented.")
bands = ["0%", "1-10%", "11-20%", "21-30%", "31-50%", "over 50%"]
head(ws5, 4, ["Discount band", "Order lines", "Gross sales", "Discount given", "Net revenue",
              "Implied cost", "Profit", "Net margin", "Loss lines %"], [15, 13, 16, 16, 16, 16, 15, 13, 13])
for i, b in enumerate(bands):
    rr = 5 + i
    ws5.cell(row=rr, column=1, value=b).font = Font(name=FONT, size=10, bold=True)
    ws5.cell(row=rr, column=2, value=f'=COUNTIF({rng("DiscountBand")},$A{rr})').number_format = INT
    for j, col in enumerate(["GrossSales", "DiscountValue", "Revenue", "ImpliedCost", "Profit"]):
        ws5.cell(row=rr, column=3 + j,
                 value=f'=SUMIF({rng("DiscountBand")},$A{rr},{rng(col)})').number_format = MONEY
    ws5.cell(row=rr, column=8, value=f"=G{rr}/E{rr}").number_format = PCT2
    ws5.cell(row=rr, column=9,
             value=f'=SUMIF({rng("DiscountBand")},$A{rr},{rng("IsLossMaking")})/B{rr}').number_format = PCT2

# ======================================================= REGIONAL ==========
ws6 = wb.create_sheet("Regional_Analysis")
title_block(ws6, "Regional performance", "By region, then every state that loses money overall.")
fin_block(ws6, 4, regs, rng("Region"), "Region", 20)
loss_states = (d.groupby("State").Profit.sum().pipe(lambda s: s[s < 0]).sort_values().index.tolist())
r6 = 4 + len(regs) + 3
ws6.cell(row=r6, column=1, value="States that lose money overall").font = Font(name=FONT, size=11, bold=True, color=ACCENT)
head(ws6, r6 + 1, ["State", "Net revenue", "Implied cost", "Profit", "Net margin", "Discount rate"],
     [22, 16, 16, 16, 14, 14])
for i, s_ in enumerate(loss_states):
    rr = r6 + 2 + i
    ws6.cell(row=rr, column=1, value=s_)
    ws6.cell(row=rr, column=2, value=f'=SUMIF({rng("State")},$A{rr},{rng("Revenue")})').number_format = MONEY
    ws6.cell(row=rr, column=3, value=f'=SUMIF({rng("State")},$A{rr},{rng("ImpliedCost")})').number_format = MONEY
    ws6.cell(row=rr, column=4, value=f'=SUMIF({rng("State")},$A{rr},{rng("Profit")})').number_format = MONEY
    ws6.cell(row=rr, column=5, value=f"=D{rr}/B{rr}").number_format = PCT2
    ws6.cell(row=rr, column=6,
             value=f'=SUMIF({rng("State")},$A{rr},{rng("DiscountValue")})'
                   f'/SUMIF({rng("State")},$A{rr},{rng("GrossSales")})').number_format = PCT2

# ======================================================= MONTHLY ===========
ws7 = wb.create_sheet("Monthly_Trends")
title_block(ws7, "Monthly trend", "48 continuous months. Growth is month over month.")
head(ws7, 4, ["Month", "Net revenue", "Implied cost", "Profit", "Net margin", "Revenue growth"],
     [12, 18, 18, 16, 14, 14])
months = sorted(d.YearMonth.unique().tolist())
for i, m in enumerate(months):
    rr = 5 + i
    ws7.cell(row=rr, column=1, value=m)
    ws7.cell(row=rr, column=2, value=f'=SUMIF({rng("YearMonth")},$A{rr},{rng("Revenue")})').number_format = MONEY
    ws7.cell(row=rr, column=3, value=f'=SUMIF({rng("YearMonth")},$A{rr},{rng("ImpliedCost")})').number_format = MONEY
    ws7.cell(row=rr, column=4, value=f'=SUMIF({rng("YearMonth")},$A{rr},{rng("Profit")})').number_format = MONEY
    ws7.cell(row=rr, column=5, value=f"=D{rr}/B{rr}").number_format = PCT2
    if i:
        ws7.cell(row=rr, column=6, value=f"=B{rr}/B{rr - 1}-1").number_format = PCT2

# ================================================== EXECUTIVE SUMMARY ======
ws8 = wb.create_sheet("Executive_Summary", 0)
title_block(ws8, "Financial Performance - executive summary",
            f"Every figure references another sheet, so the statements update with the data. Author: {AUTHOR}")
ws8.column_dimensions["A"].width = 4
ws8.column_dimensions["B"].width = 62
ws8.column_dimensions["C"].width = 16
ws8.column_dimensions["D"].width = 74
head(ws8, 4, ["#", "Finding", "Figure", "What management could do about it"])
FINDINGS = [
    ("Net margin on net revenue", f"=KPI_Summary!$B${ROWREF['Net margin']}", PCT2,
     "The business keeps about an eighth of what it sells. Protect it before chasing volume."),
    ("Cost of sales absorbs most of revenue", f"=KPI_Summary!$B${ROWREF['Cost-to-revenue ratio']}", PCT2,
     "Cost is the largest single lever; a one-point improvement is worth more than most revenue pushes."),
    ("Discount given away, as a share of list value", f"=KPI_Summary!$B${ROWREF['Discount rate']}", PCT2,
     "Nearly a fifth of list value never reaches the business. Review the discount approval policy."),
    ("Order lines that lose money", f"=KPI_Summary!$B${ROWREF['Loss-making order lines (share)']}", PCT2,
     "Almost one line in five is sold below cost. Identify them before they are quoted, not after."),
    ("Profit handed back on those lines", f"=KPI_Summary!$B${ROWREF['Profit lost on those lines']}", MONEY,
     "This is the size of the prize from fixing discounting alone."),
    ("Margin at no discount", "=Discount_Analysis!$H$5", PCT2,
     "Undiscounted business is healthy. The problem is not the product set."),
    ("Margin above a 50% discount", "=Discount_Analysis!$H$10", PCT2,
     "Deep discounts destroy more than they sell. Cap discretionary discounting."),
    ("Average order value", f"=KPI_Summary!$B${ROWREF['Average order value']}", MONEY,
     "Use with profit per order below - volume alone does not pay."),
    ("Profit per order", f"=KPI_Summary!$B${ROWREF['Profit per order']}", MONEY,
     "The real unit economics of winning an order."),
]
for i, (text, f, fmt, action) in enumerate(FINDINGS):
    rr = 5 + i
    ws8.cell(row=rr, column=1, value=i + 1).font = Font(name=FONT, size=10, bold=True, color=ACCENT)
    ws8.cell(row=rr, column=2, value=text).font = Font(name=FONT, size=10)
    v = ws8.cell(row=rr, column=3, value=f)
    v.font = Font(name=FONT, size=10, bold=True, color=ACCENT)
    v.number_format = fmt
    ws8.cell(row=rr, column=4, value=action).font = Font(name=FONT, size=9, color=MUTED)
    for col in range(1, 5):
        ws8.cell(row=rr, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        ws8.cell(row=rr, column=col).border = BORDER
    ws8.row_dimensions[rr].height = 30

# ===================================================== DATA DICTIONARY =====
ws9 = wb.create_sheet("Data_Dictionary")
title_block(ws9, "Data dictionary", f"Every field, where it came from and how it is calculated. Author: {AUTHOR}")
head(ws9, 4, ["Field", "Source", "Definition / formula"], [20, 14, 96])
DICT = [
    ("Row ID", "dataset", "Unique identifier of the order line"),
    ("Order ID", "dataset", "Order the line belongs to; one order can hold several lines"),
    ("OrderDate", "derived", "Order Date parsed from M/D/YYYY text"),
    ("Segment", "dataset", "Consumer, Corporate or Home Office - used here only as a revenue dimension"),
    ("Region / State", "dataset", "US sales region and state"),
    ("Category / Sub-Category", "dataset", "Product hierarchy: 3 categories, 17 sub-categories"),
    ("Quantity", "dataset", "Units on the line"),
    ("Discount", "dataset", "Discount rate applied, 0 to 0.80"),
    ("DiscountBand", "derived", "Fixed bands: 0%, 1-10%, 11-20%, 21-30%, 31-50%, over 50%"),
    ("Revenue", "dataset", "The dataset's Sales field. Verified to be NET of discount"),
    ("GrossSales", "derived", "Revenue / (1 - Discount) - the list value before discount"),
    ("DiscountValue", "derived", "GrossSales - Revenue - the money given away"),
    ("Profit", "dataset", "The dataset's own per-line profit figure"),
    ("ImpliedCost", "derived", "Revenue - Profit. A RESIDUAL cost of sales, NOT a reported cost line"),
    ("Net margin", "derived", "Profit / Revenue, always recomputed from sums, never averaged"),
    ("Cost-to-revenue", "derived", "ImpliedCost / Revenue"),
    ("IsLossMaking", "derived", "1 when Profit < 0"),
    ("NOT AVAILABLE", "-", "Operating expenses, overheads, payroll, tax, balance-sheet data. "
                           "No operating margin, EBITDA or accounting net profit can be computed."),
]
for i, (f_, s_, de) in enumerate(DICT):
    rr = 5 + i
    ws9.cell(row=rr, column=1, value=f_).font = Font(name=FONT, size=10, bold=True)
    ws9.cell(row=rr, column=2, value=s_).font = Font(name=FONT, size=9, color=MUTED)
    ws9.cell(row=rr, column=3, value=de).font = Font(name=FONT, size=9, color=MUTED)
    for col in range(1, 4):
        ws9.cell(row=rr, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        ws9.cell(row=rr, column=col).border = BORDER

wb.save(OUT)
print(f"  saved {OUT.relative_to(ROOT)}")
print(f"  sheets: {wb.sheetnames}")
